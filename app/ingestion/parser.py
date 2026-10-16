"""Extract plain text from PDF, DOCX and TXT uploads."""
from __future__ import annotations

import io
from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".txt", ".md"}


class ParseError(ValueError):
    pass


def parse_pdf(data: bytes) -> str:
    try:
        import fitz  # PyMuPDF
    except ImportError as e:  # pragma: no cover
        raise ParseError("PyMuPDF is not installed") from e
    try:
        with fitz.open(stream=data, filetype="pdf") as doc:
            return "\n".join(page.get_text("text") for page in doc)
    except Exception as e:
        raise ParseError(f"Could not read PDF: {e}") from e


def parse_docx(data: bytes) -> str:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception as e:
        raise ParseError(f"Could not read DOCX: {e}") from e
    lines = [p.text for p in doc.paragraphs]
    # resumes often keep skills/experience inside tables
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(dict.fromkeys(cells)))
    return "\n".join(lines)


def parse_document(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED:
        raise ParseError(f"Unsupported file type '{ext}'. Use one of: {', '.join(sorted(SUPPORTED))}")
    if ext == ".pdf":
        text = parse_pdf(data)
    elif ext == ".docx":
        text = parse_docx(data)
    else:
        text = data.decode("utf-8", errors="ignore")
    text = clean_text(text)
    if len(text.split()) < 20:
        raise ParseError("Extracted text is too short. Scanned PDFs need OCR, which is not supported yet.")
    return text


def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\u00a0", " ")
    lines = [" ".join(line.split()) for line in text.split("\n")]
    out, blank = [], 0
    for line in lines:
        blank = blank + 1 if not line else 0
        if blank <= 1:
            out.append(line)
    return "\n".join(out).strip()
